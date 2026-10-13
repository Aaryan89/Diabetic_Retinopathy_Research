"""
Convert Research Paper Markdown to Authentic IEEE Conference Format DOCX and PDF
Based on conference-template-letter.docx specification:
- US Letter, exact IEEE margins (0.75" top, 1.0" bottom, 0.62" left/right)
- 4-column author block with Name, PRN, Affiliation, and Email
- 2-column layout for body text and figures
- Full-width tables spanning across both columns (consuming all horizontal space)
- Native OMML equation typesetting (123+ inline math, 14 block equations)
- Exact IEEE heading hierarchy, table styling, and hanging-indent references
"""

import os
import re
import zipfile
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import latex2mathml.converter
import lxml.etree as ET

# Office MathML to OMML XSLT stylesheet
XSL_PATH = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
if os.path.exists(XSL_PATH):
    _xslt_doc = ET.parse(XSL_PATH)
    _transform = ET.XSLT(_xslt_doc)
else:
    _transform = None
    print(f"[WARNING] MML2OMML.XSL not found at: {XSL_PATH}")

def latex_to_omml(latex_str, is_block=False):
    """Converts a LaTeX math string into a docx.oxml OMML element."""
    if _transform is None:
        return None
    try:
        clean_latex = latex_str.strip()
        clean_latex = re.sub(r'\\implies', r'\\Rightarrow', clean_latex)
        mathml = latex2mathml.converter.convert(clean_latex)
        dom = ET.fromstring(mathml.encode('utf-8'))
        omml_tree = _transform(dom)
        root = omml_tree.getroot()
        omml_xml = ET.tostring(root, encoding='utf-8').decode('utf-8')
        
        if is_block:
            wrapped_xml = f'<m:oMathPara {nsdecls("m")}>{omml_xml}</m:oMathPara>'
            return parse_xml(wrapped_xml)
        else:
            return parse_xml(omml_xml)
    except Exception:
        return None

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_width(cell, width_in_inches):
    """Sets explicit width on cell to prevent wrapping."""
    width_dxa = int(width_in_inches * 1440)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_table_borders(table):
    """Adds subtle professional IEEE borders to table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="2B5C8F"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="333333"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_section_columns(section, num_cols, space_dxa=360):
    """Configures the number of columns and gutter spacing on a Word section."""
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(num_cols))
    if num_cols > 1:
        cols.set(qn('w:space'), str(space_dxa))
    elif qn('w:space') in cols.attrib:
        del cols.attrib[qn('w:space')]

def _render_token(paragraph, token, base_font_size, is_bold, is_italic, base_color):
    """Renders a single token or converts inline math to OMML."""
    if token.startswith('$') and token.endswith('$') and len(token) > 2:
        latex_content = token[1:-1].strip()
        if re.match(r'^[0-9,.\s%+\-]+$', latex_content):
            run = paragraph.add_run(latex_content)
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.bold = is_bold
            run.font.italic = is_italic
            run.font.color.rgb = base_color
            return
        
        omml_elem = latex_to_omml(latex_content, is_block=False)
        if omml_elem is not None:
            paragraph._p.append(omml_elem)
        else:
            clean_text = latex_content.replace('\\', '')
            run = paragraph.add_run(clean_text)
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.bold = is_bold
            run.font.italic = True
        return

    run = paragraph.add_run(token)
    run.font.name = 'Times New Roman'
    run.font.size = base_font_size
    run.font.bold = is_bold
    run.font.italic = is_italic
    run.font.color.rgb = base_color

def add_formatted_runs(paragraph, text, base_font_size=Pt(9.5), base_italic=False, base_bold=False, base_color=RGBColor(30, 30, 30)):
    """Parses markdown bold, italic, code, and inline LaTeX math."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$[^\$\n]+?\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('$') and token.endswith('$') and len(token) > 2:
            _render_token(paragraph, token, base_font_size, base_bold, base_italic, base_color)
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            if '$' in inner:
                inner_tokens = re.split(r'(\$[^\$\n]+?\$)', inner)
                for it in inner_tokens:
                    if it:
                        _render_token(paragraph, it, base_font_size, is_bold=True, is_italic=base_italic, base_color=base_color)
            else:
                _render_token(paragraph, inner, base_font_size, is_bold=True, is_italic=base_italic, base_color=base_color)
            continue
        if token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            if '$' in inner:
                inner_tokens = re.split(r'(\$[^\$\n]+?\$)', inner)
                for it in inner_tokens:
                    if it:
                        _render_token(paragraph, it, base_font_size, is_bold=base_bold, is_italic=True, base_color=base_color)
            else:
                _render_token(paragraph, inner, base_font_size, is_bold=base_bold, is_italic=True, base_color=base_color)
            continue
        if token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(60, 60, 60)
            continue
        if '$' in token:
            inner_tokens = re.split(r'(\$[^\$\n]+?\$)', token)
            for it in inner_tokens:
                if it:
                    _render_token(paragraph, it, base_font_size, is_bold=base_bold, is_italic=base_italic, base_color=base_color)
        else:
            _render_token(paragraph, token, base_font_size, is_bold=base_bold, is_italic=base_italic, base_color=base_color)

def get_full_width_table_columns(header_row, num_cols):
    """
    Computes precise column widths summing to 7.25 inches (full IEEE printable page width)
    to guarantee zero mid-word wrapping.
    """
    hdr_text = " ".join(header_row).lower()
    
    # 1. Table 1: Literature Survey (7 cols)
    if ('architecture' in hdr_text or 'relation to our work' in hdr_text or 'identified architectural gap' in hdr_text) and num_cols == 7:
        return [1.10, 1.25, 0.85, 0.70, 0.85, 1.25, 1.25], Pt(7.5), (60, 60)

    # 2. Table 2: Hyperparameters / Training Details (6 cols)
    if ('hyperparameter' in hdr_text or 'training configuration' in hdr_text) and num_cols == 6:
        return [1.75, 1.10, 1.10, 1.10, 1.10, 1.10], Pt(7.5), (55, 55)

    # 3. Table 3: Model Comparison (9 cols)
    if 'qwk (primary)' in hdr_text and 'exact match' in hdr_text:
        return [1.55, 0.90, 0.55, 0.80, 0.85, 0.65, 0.65, 0.65, 0.65], Pt(7.2), (50, 50)

    # 4. Table 4: Per-Class Diagnostic Performance (5 cols)
    if 'recall / sensitivity' in hdr_text or 'class grade' in hdr_text:
        return [1.25, 1.80, 1.40, 1.40, 1.40], Pt(8.0), (70, 70)

    # 5. Table 5: Ensemble Configurations (7 cols)
    if ('model configuration' in hdr_text or 'ensemble configuration' in hdr_text) and 'catastrophic' in hdr_text:
        return [1.75, 0.90, 0.90, 0.85, 0.95, 0.95, 0.95], Pt(7.5), (60, 60)

    # 6. Table 6: Post-Hoc Logit Adjustment Temperature Sweep (6 cols)
    if ('temperature' in hdr_text or 'logit adjustment' in hdr_text) and num_cols == 6:
        return [1.85, 0.95, 0.95, 0.90, 1.30, 1.30], Pt(8.0), (70, 70)

    # 7. Table 7: Clinical Referral Triage (6 cols)
    if 'under-triage' in hdr_text and 'over-referral' in hdr_text:
        return [1.85, 1.10, 0.95, 1.10, 0.95, 1.30], Pt(7.5), (55, 55)

    # Default fallback: uniform distribution across 7.25 inches
    w = round(7.25 / num_cols, 2)
    return [w] * num_cols, Pt(8.0), (60, 60)

def render_full_width_table(doc, table_lines):
    """
    Renders a full-width table spanning across both columns (consuming all horizontal space)
    by switching to 1-column section, rendering the table across 7.25 inches, and switching back.
    """
    if len(table_lines) < 2:
        return

    rows_data = []
    for line in table_lines:
        parts = [p.strip() for p in line.strip('|').split('|')]
        if all(re.match(r'^:?-+:?$', p) for p in parts):
            continue
        rows_data.append(parts)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    header_row = rows_data[0]
    hdr_text = " ".join(header_row).lower()
    col_widths, font_size, (pad_tb, pad_lr) = get_full_width_table_columns(header_row, num_cols)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    for r_idx, row_data in enumerate(rows_data):
        row = table.rows[r_idx]
        is_header = (r_idx == 0)

        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            w_in = col_widths[c_idx] if c_idx < len(col_widths) else 1.0
            set_cell_width(cell, w_in)
            set_cell_margins(cell, top=pad_tb, bottom=pad_tb, left=pad_lr, right=pad_lr)

            text = row_data[c_idx] if c_idx < len(row_data) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or c_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

            if is_header:
                set_cell_background(cell, "2B5C8F")
                add_formatted_runs(p, text, base_font_size=font_size, base_bold=True, base_color=RGBColor(255, 255, 255))
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F7F9FC")
                else:
                    set_cell_background(cell, "FFFFFF")
                add_formatted_runs(p, text, base_font_size=font_size, base_color=RGBColor(30, 30, 30))

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(2)

def build_author_block(doc):
    """Constructs the IEEE 4-author block in a 1-row, 4-column borderless table."""
    authors = [
        {
            "name": "Nayna Sharma",
            "prn": "1262241524",
            "dept": "Dept. of Computer Science & Engg.",
            "inst": "MIT World Peace University",
            "city": "Pune, India",
            "email": "nayna.sharma@mitwpu.edu.in"
        },
        {
            "name": "Harshad Pardhi",
            "prn": "1262241534",
            "dept": "Dept. of Computer Science & Engg.",
            "inst": "MIT World Peace University",
            "city": "Pune, India",
            "email": "harshad.pardhi@mitwpu.edu.in"
        },
        {
            "name": "Aaryan Kumbhare",
            "prn": "1262241729",
            "dept": "Dept. of Computer Science & Engg.",
            "inst": "MIT World Peace University",
            "city": "Pune, India",
            "email": "aaryan.kumbhare@mitwpu.edu.in"
        },
        {
            "name": "Anurag Harapanahalli",
            "prn": "1262241775",
            "dept": "Dept. of Computer Science & Engg.",
            "inst": "MIT World Peace University",
            "city": "Pune, India",
            "email": "anurag.harapanahalli@mitwpu.edu.in"
        }
    ]

    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Borderless
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    cell_width = 7.25 / 4.0 # 1.8125 inches per author
    row = tbl.rows[0]
    for idx, auth in enumerate(authors):
        cell = row.cells[idx]
        set_cell_width(cell, cell_width)
        set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        
        # Line 1: Name (Bold, 10 pt)
        r_name = p.add_run(auth["name"] + "\n")
        r_name.font.name = 'Times New Roman'
        r_name.font.size = Pt(10)
        r_name.font.bold = True
        
        # Line 2: PRN
        r_prn = p.add_run(f"PRN: {auth['prn']}\n")
        r_prn.font.name = 'Times New Roman'
        r_prn.font.size = Pt(8.5)
        
        # Line 3: Department
        r_dept = p.add_run(f"{auth['dept']}\n")
        r_dept.font.name = 'Times New Roman'
        r_dept.font.size = Pt(8.5)
        r_dept.font.italic = True
        
        # Line 4: Institution
        r_inst = p.add_run(f"{auth['inst']}\n")
        r_inst.font.name = 'Times New Roman'
        r_inst.font.size = Pt(8.5)
        r_inst.font.italic = True
        
        # Line 5: City, Country
        r_city = p.add_run(f"{auth['city']}\n")
        r_city.font.name = 'Times New Roman'
        r_city.font.size = Pt(8.5)
        
        # Line 6: Email
        r_email = p.add_run(auth["email"])
        r_email.font.name = 'Times New Roman'
        r_email.font.size = Pt(8.5)

    # Mentor, Course, and Academic Year banner below authors
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(10)
    p_meta.paragraph_format.line_spacing = 1.05
    run_meta = p_meta.add_run("Course: Artificial Intelligence & Expert Systems Laboratory (CSE30070)  •  Project Mentor: Prof. Pramod Mali  •  Academic Year: 2026–2027")
    run_meta.font.name = 'Times New Roman'
    run_meta.font.size = Pt(8.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(70, 70, 70)

def convert_main_sub_to_docx(md_path, docx_path, base_image_dir="."):
    """Converts main_sub.md into IEEE conference formatted docx with 2-column body and full-width tables."""
    doc = docx.Document()

    # Section 0: Title and Authors (US Letter, 1 column)
    s0 = doc.sections[0]
    s0.page_width = Inches(8.5)
    s0.page_height = Inches(11.0)
    s0.top_margin = Inches(0.75)     # 54 pt
    s0.bottom_margin = Inches(1.0)    # 72 pt
    s0.left_margin = Inches(0.62)     # 44.65 pt
    s0.right_margin = Inches(0.62)    # 44.65 pt
    set_section_columns(s0, 1)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.15
    run_title = p_title.add_run("Ordinal Neural Networks for Diabetic Retinopathy Staging: Enforcing Monotonic Rank Consistency for Safety-Critical Clinical Decision Support")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 30, 60)

    # 4-Author Block
    build_author_block(doc)

    # Switch to 2 columns for Abstract and Body text
    s_body = doc.add_section(WD_SECTION_START.CONTINUOUS)
    set_section_columns(s_body, 2, space_dxa=360)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Skip lines up to ## 1. ABSTRACT
    start_idx = 0
    for idx, line in enumerate(lines):
        if '## 1. ABSTRACT' in line:
            start_idx = idx
            break

    in_table = False
    table_lines = []
    figure_counter = 1
    table_header_pending = None

    i = start_idx
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle empty line
        if not stripped:
            if in_table and table_lines:
                # Switch to 1 column for table
                s_tbl = doc.add_section(WD_SECTION_START.CONTINUOUS)
                set_section_columns(s_tbl, 1)
                
                if table_header_pending:
                    p_th = doc.add_paragraph()
                    p_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_th.paragraph_format.space_before = Pt(6)
                    p_th.paragraph_format.space_after = Pt(2)
                    p_th.paragraph_format.keep_with_next = True
                    add_formatted_runs(p_th, table_header_pending, base_font_size=Pt(9), base_bold=True, base_color=RGBColor(20, 40, 80))
                    table_header_pending = None

                render_full_width_table(doc, table_lines)
                table_lines = []
                in_table = False

                # Switch back to 2 columns for text
                s_back = doc.add_section(WD_SECTION_START.CONTINUOUS)
                set_section_columns(s_back, 2, space_dxa=360)
            i += 1
            continue

        # Check for table rows
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # Table ended
            s_tbl = doc.add_section(WD_SECTION_START.CONTINUOUS)
            set_section_columns(s_tbl, 1)
            
            if table_header_pending:
                p_th = doc.add_paragraph()
                p_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_th.paragraph_format.space_before = Pt(6)
                p_th.paragraph_format.space_after = Pt(2)
                p_th.paragraph_format.keep_with_next = True
                add_formatted_runs(p_th, table_header_pending, base_font_size=Pt(9), base_bold=True, base_color=RGBColor(20, 40, 80))
                table_header_pending = None

            render_full_width_table(doc, table_lines)
            table_lines = []
            in_table = False

            s_back = doc.add_section(WD_SECTION_START.CONTINUOUS)
            set_section_columns(s_back, 2, space_dxa=360)

        # Table caption header (#### TABLE X: ...)
        if stripped.startswith('#### ') and 'TABLE' in stripped.upper():
            table_header_pending = stripped[5:].strip()
            i += 1
            continue

        # Table note paragraph (*Note: ...*) immediately following table in 1 column
        if stripped.startswith('*Note:') or stripped.startswith('*(Note:'):
            p_note = doc.add_paragraph()
            p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_note.paragraph_format.space_before = Pt(1)
            p_note.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p_note, stripped, base_font_size=Pt(8), base_italic=True, base_color=RGBColor(60, 60, 60))
            i += 1
            continue

        # Block math ($$...$$)
        if stripped.startswith('$$'):
            math_lines = []
            if stripped.endswith('$$') and len(stripped) > 4:
                math_content = stripped[2:-2].strip()
                i += 1
            else:
                math_lines.append(stripped[2:])
                i += 1
                while i < len(lines):
                    curr = lines[i].rstrip('\r\n').strip()
                    if curr.endswith('$$'):
                        math_lines.append(curr[:-2])
                        i += 1
                        break
                    else:
                        math_lines.append(curr)
                        i += 1
                math_content = "\n".join(math_lines).strip()

            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(4)
            p_eq.paragraph_format.space_after = Pt(4)
            
            omml_block = latex_to_omml(math_content, is_block=True)
            if omml_block is not None:
                p_eq._p.append(omml_block)
            else:
                p_eq.add_run(math_content)
            continue

        # Image tag: ![Caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1).strip()
            img_rel_path = img_match.group(2).strip()
            img_full_path = os.path.normpath(os.path.join(base_image_dir, img_rel_path))
            
            if not os.path.exists(img_full_path):
                if os.path.exists(img_rel_path):
                    img_full_path = img_rel_path
                elif os.path.exists(os.path.join("submission", "results", os.path.basename(img_rel_path))):
                    img_full_path = os.path.join("submission", "results", os.path.basename(img_rel_path))

            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(4)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                
                # Image scaled to 3.4 inches to fit cleanly inside 3.5 in column
                run_img.add_picture(img_full_path, width=Inches(3.40))

                clean_cap = re.sub(r'^(Fig\.|Figure)\s*\d*[:.]?\s*', '', caption, flags=re.IGNORECASE).strip()
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(5)
                add_formatted_runs(p_cap, f"Fig. {figure_counter}. {clean_cap}", base_font_size=Pt(8.0), base_italic=True, base_color=RGBColor(60, 60, 60))
                figure_counter += 1
            else:
                p_err = doc.add_paragraph(f"[Image not found: {img_rel_path}]")
                p_err.runs[0].font.italic = True
            i += 1
            continue

        # Horizontal rule (---)
        if stripped == '---':
            i += 1
            continue

        # Heading 1 (## 1. ABSTRACT, ## 2. LITERATURE SURVEY, etc.)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            
            # Roman numeral mapping for IEEE Section Heads
            num_match = re.match(r'^(\d+)\.\s*(.*)', h_text)
            if num_match:
                sec_num = int(num_match.group(1))
                sec_title = num_match.group(2).strip().upper()
                roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
                roman = roman_numerals[sec_num - 1] if sec_num <= len(roman_numerals) else str(sec_num)
                ieee_h1 = f"{roman}.  {sec_title}"
            else:
                ieee_h1 = h_text.upper()

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            run = p.add_run(ieee_h1)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 30, 60)
            i += 1
            continue

        # Heading 2 (### 2.1 Analysis of Existing Literature, etc.)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            
            # Map "X.Y Title" to "A. Title", "B. Title"
            sub_match = re.match(r'^\d+\.(\d+)\s*(.*)', h_text)
            if sub_match:
                sub_idx = int(sub_match.group(1))
                sub_title = sub_match.group(2).strip()
                letter = chr(ord('A') + sub_idx - 1) if sub_idx <= 26 else str(sub_idx)
                ieee_h2 = f"{letter}.  {sub_title}"
            else:
                ieee_h2 = h_text

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(ieee_h2)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(30, 40, 70)
            i += 1
            continue

        # Heading 4 / Subheading (##### 1) Architectural..., etc.)
        if stripped.startswith('##### '):
            h_text = stripped[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            add_formatted_runs(p, h_text, base_font_size=Pt(9.0), base_bold=True, base_color=RGBColor(40, 50, 70))
            i += 1
            continue

        # Heading 3 (#### 1. Khalifa et al., etc.)
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            add_formatted_runs(p, h_text, base_font_size=Pt(9.0), base_bold=True, base_italic=True, base_color=RGBColor(40, 50, 70))
            i += 1
            continue

        # References ([1] Authors...)
        ref_match = re.match(r'^\[(\d+)\]\s+(.*)', stripped)
        if ref_match:
            ref_num = ref_match.group(1)
            ref_body = ref_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            run_num = p.add_run(f"[{ref_num}]\t")
            run_num.font.name = 'Times New Roman'
            run_num.font.size = Pt(8)
            run_num.font.bold = True
            
            add_formatted_runs(p, ref_body, base_font_size=Pt(8))
            i += 1
            continue

        # List item (- or * or numbered 1.)
        bullet_match = re.match(r'^([-*]|\d+\.)\s+(.*)', stripped)
        if bullet_match:
            marker = bullet_match.group(1)
            item_text = bullet_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.0)
            p.paragraph_format.line_spacing = 1.05
            
            run_bullet = p.add_run("• " if marker in ['-', '*'] else f"{marker} ")
            run_bullet.font.name = 'Times New Roman'
            run_bullet.font.size = Pt(9.5)
            run_bullet.font.bold = True
            
            add_formatted_runs(p, item_text, base_font_size=Pt(9.5))
            i += 1
            continue

        # Standard Body Paragraph (including Abstract lead-in)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.05

        # Check if this paragraph is the Abstract body text
        if stripped.startswith("Automated grading of Diabetic Retinopathy"):
            # Prefix with IEEE Abstract— lead-in
            run_abs = p.add_run("Abstract—")
            run_abs.font.name = 'Times New Roman'
            run_abs.font.size = Pt(9.0)
            run_abs.font.bold = True
            run_abs.font.italic = True
            add_formatted_runs(p, stripped, base_font_size=Pt(9.0))
        elif stripped.startswith("In this work, we evaluate rank-consistent"):
            # Second paragraph of abstract
            add_formatted_runs(p, stripped, base_font_size=Pt(9.0))
        elif stripped.startswith("Our featured model, **Variant E: 3-Seed Ensemble"):
            # Third paragraph of abstract + append IEEE Index Terms
            add_formatted_runs(p, stripped, base_font_size=Pt(9.0))
            
            # Keywords paragraph
            p_kw = doc.add_paragraph()
            p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_kw.paragraph_format.space_before = Pt(3)
            p_kw.paragraph_format.space_after = Pt(6)
            p_kw.paragraph_format.line_spacing = 1.05
            run_kw = p_kw.add_run("Keywords—")
            run_kw.font.name = 'Times New Roman'
            run_kw.font.size = Pt(9.0)
            run_kw.font.bold = True
            run_kw.font.italic = True
            run_kw_body = p_kw.add_run("Diabetic retinopathy grading, ordinal regression, conditional ordinal regression (CORN), Quadratic Weighted Kappa (QWK), clinical triage, test-time augmentation, multi-seed ensemble, EfficientNet-B0.")
            run_kw_body.font.name = 'Times New Roman'
            run_kw_body.font.size = Pt(9.0)
        else:
            add_formatted_runs(p, stripped, base_font_size=Pt(9.5))
        i += 1

    # Flush trailing table if needed
    if in_table and table_lines:
        s_tbl = doc.add_section(WD_SECTION_START.CONTINUOUS)
        set_section_columns(s_tbl, 1)
        if table_header_pending:
            p_th = doc.add_paragraph()
            p_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_th.paragraph_format.space_before = Pt(6)
            p_th.paragraph_format.space_after = Pt(2)
            p_th.paragraph_format.keep_with_next = True
            add_formatted_runs(p_th, table_header_pending, base_font_size=Pt(9), base_bold=True, base_color=RGBColor(20, 40, 80))
        render_full_width_table(doc, table_lines)

    doc.save(docx_path)
    print(f"[SUCCESS] Exported IEEE Conference formatted DOCX to: {docx_path}")

def convert_docx_to_pdf(docx_path, pdf_path):
    """Converts DOCX to PDF using Word COM automation."""
    import win32com.client as win32
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)
    
    if os.path.exists(abs_pdf):
        try:
            os.remove(abs_pdf)
        except Exception:
            pass

    word = win32.gencache.EnsureDispatch('Word.Application')
    word.Visible = False
    doc = word.Documents.Open(abs_docx)
    doc.ExportAsFixedFormat(abs_pdf, 17) # 17 = wdExportFormatPDF
    doc.Close(False)
    word.Quit()
    print(f"[SUCCESS] Exported PDF to: {abs_pdf}")

if __name__ == "__main__":
    import sys
    md_in = sys.argv[1] if len(sys.argv) > 1 else "main_sub.md"
    docx_out = sys.argv[2] if len(sys.argv) > 2 else "main_sub.docx"
    pdf_out = sys.argv[3] if len(sys.argv) > 3 else "main_sub.pdf"
    
    convert_main_sub_to_docx(md_in, docx_out)
    convert_docx_to_pdf(docx_out, pdf_out)
