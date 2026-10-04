"""
Markdown to IEEE-Formatted DOCX Converter with Native OMML Equation Typesetting
Transforms research paper markdown documents into beautifully styled Word documents (.docx)
with full support for native Word OMML equations (LaTeX math conversion), custom IEEE table layouts,
column width calculations, embedded figures with sequential numbering, and hanging-indent references.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import latex2mathml.converter
import lxml.etree as ET

# Office MathML to OMML XSLT stylesheet
XSL_PATH = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
if os.path.exists(XSL_PATH):
    _xslt_doc = ET.parse(XSL_PATH)
    _transform = ET.XSLT(_xslt_doc)
else:
    _transform = None
    print(f"[WARNING] MML2OMML.XSL not found at: {XSL_PATH}")

def latex_to_omml(latex_str, is_block=False):
    """Converts a LaTeX math string into a docx.oxml OMML element."""
    if _transform is None:
        return None
    try:
        clean_latex = latex_str.strip()
        # Common entity replacements
        clean_latex = re.sub(r'\\implies', r'\\Rightarrow', clean_latex)
        # Convert to MathML
        mathml = latex2mathml.converter.convert(clean_latex)
        dom = ET.fromstring(mathml.encode('utf-8'))
        omml_tree = _transform(dom)
        root = omml_tree.getroot()
        omml_xml = ET.tostring(root, encoding='utf-8').decode('utf-8')
        
        if is_block:
            # Wrap in oMathPara for display math
            wrapped_xml = f'<m:oMathPara {nsdecls("m")}>{omml_xml}</m:oMathPara>'
            return parse_xml(wrapped_xml)
        else:
            return parse_xml(omml_xml)
    except Exception as e:
        # Fallback to plain text on rare syntax edge cases
        return None

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_width(cell, width_in_inches):
    """Sets explicit width on cell to prevent wrapping."""
    width_dxa = int(width_in_inches * 1440)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_table_borders(table):
    """Adds subtle professional borders to table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="2B5C8F"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="333333"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def _render_token(paragraph, token, base_font_size, is_bold, is_italic, base_color):
    """Renders a single token or converts inline math to OMML."""
    if token.startswith('$') and token.endswith('$') and len(token) > 2:
        latex_content = token[1:-1].strip()
        # If it's a simple number or percentage, add as clean run
        if re.match(r'^[0-9,.\s%+\-]+$', latex_content):
            run = paragraph.add_run(latex_content)
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.bold = is_bold
            run.font.italic = is_italic
            run.font.color.rgb = base_color
            return
        
        # Convert LaTeX to native OMML
        omml_elem = latex_to_omml(latex_content, is_block=False)
        if omml_elem is not None:
            paragraph._p.append(omml_elem)
        else:
            clean_text = latex_content.replace('\\', '')
            run = paragraph.add_run(clean_text)
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.bold = is_bold
            run.font.italic = True
        return

    # Check for action required / pending indicator
    if '[ACTION REQUIRED:' in token or '[Pending' in token:
        run = paragraph.add_run(token)
        run.font.name = 'Times New Roman'
        run.font.size = base_font_size
        run.font.bold = True
        run.font.color.rgb = RGBColor(190, 30, 20) # Deep crimson for visibility
        return

    run = paragraph.add_run(token)
    run.font.name = 'Times New Roman'
    run.font.size = base_font_size
    run.font.bold = is_bold
    run.font.italic = is_italic
    run.font.color.rgb = base_color

def add_formatted_runs(paragraph, text, base_font_size=Pt(10), base_italic=False, base_bold=False, base_color=RGBColor(30, 30, 30)):
    """
    Parses markdown bold, italic, inline code, and inline LaTeX math ($...$)
    converting math into genuine OMML native Word equations, including nested math inside bold/italic spans.
    """
    # Pattern to match bold (**), italic (*), code (`), or standalone math ($...$)
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$[^\$\n]+?\$)', text)
    
    for token in tokens:
        if not token:
            continue
        
        # 1. Standalone inline math ($...$)
        if token.startswith('$') and token.endswith('$') and len(token) > 2:
            _render_token(paragraph, token, base_font_size, base_bold, base_italic, base_color)
            continue

        # 2. Bold span (**...**)
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            if '$' in inner:
                inner_tokens = re.split(r'(\$[^\$\n]+?\$)', inner)
                for it in inner_tokens:
                    if it:
                        _render_token(paragraph, it, base_font_size, is_bold=True, is_italic=base_italic, base_color=base_color)
            else:
                _render_token(paragraph, inner, base_font_size, is_bold=True, is_italic=base_italic, base_color=base_color)
            continue

        # 3. Italic span (*...*)
        if token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            if '$' in inner:
                inner_tokens = re.split(r'(\$[^\$\n]+?\$)', inner)
                for it in inner_tokens:
                    if it:
                        _render_token(paragraph, it, base_font_size, is_bold=base_bold, is_italic=True, base_color=base_color)
            else:
                _render_token(paragraph, inner, base_font_size, is_bold=base_bold, is_italic=True, base_color=base_color)
            continue

        # 4. Inline code (`...`)
        if token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(60, 60, 60)
            continue

        # 5. Plain text (handle any residual math)
        if '$' in token:
            inner_tokens = re.split(r'(\$[^\$\n]+?\$)', token)
            for it in inner_tokens:
                if it:
                    _render_token(paragraph, it, base_font_size, is_bold=base_bold, is_italic=base_italic, base_color=base_color)
        else:
            _render_token(paragraph, token, base_font_size, is_bold=base_bold, is_italic=base_italic, base_color=base_color)

def convert_markdown_to_docx(md_path, docx_path, base_image_dir="."):
    """Converts markdown file into an IEEE-formatted DOCX with native OMML equations."""
    doc = docx.Document()

    # Set page margins (0.75 inch for academic paper feel, 7.0 inch printable width)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    figure_counter = 1
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle empty lines
        if not stripped:
            if in_table and table_lines:
                render_table(doc, table_lines)
                table_lines = []
                in_table = False
            i += 1
            continue

        # Check for table rows
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # Table ended
            render_table(doc, table_lines)
            table_lines = []
            in_table = False

        # Check for Block Math ($$...$$)
        if stripped.startswith('$$'):
            # Collect block math lines
            math_lines = []
            if stripped.endswith('$$') and len(stripped) > 4:
                # Single-line block math: $$...$$
                math_content = stripped[2:-2].strip()
                i += 1
            else:
                # Multi-line block math
                math_lines.append(stripped[2:])
                i += 1
                while i < len(lines):
                    curr = lines[i].rstrip('\r\n').strip()
                    if curr.endswith('$$'):
                        math_lines.append(curr[:-2])
                        i += 1
                        break
                    else:
                        math_lines.append(curr)
                        i += 1
                math_content = "\n".join(math_lines).strip()
            
            # Create centered equation paragraph
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(6)
            p_eq.paragraph_format.space_after = Pt(6)
            
            omml_block = latex_to_omml(math_content, is_block=True)
            if omml_block is not None:
                p_eq._p.append(omml_block)
            else:
                # Fallback run
                p_eq.add_run(math_content)
            continue

        # Image tag: ![Caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1).strip()
            img_rel_path = img_match.group(2).strip()
            img_full_path = os.path.normpath(os.path.join(base_image_dir, img_rel_path))
            
            if not os.path.exists(img_full_path):
                if os.path.exists(img_rel_path):
                    img_full_path = img_rel_path
                elif os.path.exists(os.path.join("submission", "results", os.path.basename(img_rel_path))):
                    img_full_path = os.path.join("submission", "results", os.path.basename(img_rel_path))

            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(3)
                run_img = p_img.add_run()
                
                # Proportional image scaling
                img_name = os.path.basename(img_full_path).lower()
                if "uml_activity" in img_name:
                    pic_width = Inches(5.0)
                elif "class_distribution" in img_name or "error_distance" in img_name:
                    pic_width = Inches(4.8)
                elif "sample_images" in img_name:
                    pic_width = Inches(5.2)
                elif "all_confusion" in img_name or "confusion_matrix" in img_name:
                    pic_width = Inches(5.4)
                elif "system_architecture" in img_name:
                    pic_width = Inches(5.3)
                elif "uml_class" in img_name:
                    pic_width = Inches(5.4)
                else:
                    pic_width = Inches(5.0)
                    
                run_img.add_picture(img_full_path, width=pic_width)

                # Sequential caption without duplicate prefix
                clean_cap = re.sub(r'^(Fig\.|Figure)\s*\d*[:.]?\s*', '', caption, flags=re.IGNORECASE).strip()
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(10)
                add_formatted_runs(p_cap, f"Fig. {figure_counter}. {clean_cap}", base_font_size=Pt(9), base_italic=True, base_color=RGBColor(60, 60, 60))
                figure_counter += 1
            else:
                p_err = doc.add_paragraph(f"[Image not found: {img_rel_path}]")
                p_err.runs[0].font.italic = True
            i += 1
            continue

        # Document Title (# Title)
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 30, 60)
            i += 1
            continue

        # Authors / Affiliation Block
        if stripped.startswith('> '):
            clean_meta = stripped.lstrip('> ').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, clean_meta, base_font_size=Pt(9.5), base_color=RGBColor(60, 60, 60))
            i += 1
            continue

        # Heading 1 (## 1. ABSTRACT, ## 2. LITERATURE SURVEY, etc.)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            if 'ABSTRACT' in h_text.upper():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(h_text.upper())
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(20, 20, 20)
            else:
                run = p.add_run(h_text)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(20, 40, 80)
            i += 1
            continue

        # Heading 2 (### 2.1 Detailed Analysis, etc.)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(40, 50, 70)
            i += 1
            continue

        # Heading 3 (#### 1. Khalifa et al., #### TABLE 1, etc.)
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            
            # If it's a table header label like "#### TABLE 1: ...", center and style prominently
            if h_text.upper().startswith('TABLE'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_formatted_runs(p, h_text, base_font_size=Pt(9.5), base_bold=True, base_color=RGBColor(20, 40, 80))
            else:
                add_formatted_runs(p, h_text, base_font_size=Pt(10), base_bold=True, base_color=RGBColor(50, 50, 50))
            i += 1
            continue

        # Horizontal rule (---)
        if stripped == '---':
            i += 1
            continue

        # References in Section 10 ([1] Authors, "Title"...)
        ref_match = re.match(r'^\[(\d+)\]\s+(.*)', stripped)
        if ref_match:
            ref_num = ref_match.group(1)
            ref_body = ref_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            run_num = p.add_run(f"[{ref_num}]\t")
            run_num.font.name = 'Times New Roman'
            run_num.font.size = Pt(9.5)
            run_num.font.bold = True
            
            add_formatted_runs(p, ref_body, base_font_size=Pt(9.5))
            i += 1
            continue

        # List item (- or * or numbered 1.)
        bullet_match = re.match(r'^([-*]|\d+\.)\s+(.*)', stripped)
        if bullet_match:
            marker = bullet_match.group(1)
            item_text = bullet_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            run_bullet = p.add_run("• " if marker in ['-', '*'] else f"{marker} ")
            run_bullet.font.name = 'Times New Roman'
            run_bullet.font.size = Pt(10)
            run_bullet.font.bold = True
            
            add_formatted_runs(p, item_text, base_font_size=Pt(10))
            i += 1
            continue

        # Standard Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4.5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped, base_font_size=Pt(10))
        i += 1

    # Clean up trailing table if file ends with table
    if in_table and table_lines:
        render_table(doc, table_lines)

    doc.save(docx_path)
    print(f"[SUCCESS] Exported IEEE-formatted Word document with native OMML math to: {docx_path}")

def get_table_column_widths(header_row, num_cols):
    """
    Computes precise column widths in inches summing to 7.0 inches (printable width)
    based on table type and column count to guarantee zero mid-word line breaks.
    """
    hdr_text = " ".join(header_row).lower()
    
    # 1. Team Table
    if 'student name' in hdr_text or 'prn' in hdr_text:
        if num_cols == 3:
            return [0.90, 2.50, 3.60], Pt(9.5), (90, 90)
        return [0.70, 2.70, 1.80, 1.80], Pt(8.5), (80, 80)

    # 2. Table 1: Literature Survey (7 columns)
    if 'architecture' in hdr_text and 'decisive' in hdr_text:
        return [1.05, 1.25, 0.80, 0.65, 0.80, 1.20, 1.25], Pt(7.5), (60, 60)

    # 3. Table 2: Model Comparison (9 columns)
    if 'qwk (primary)' in hdr_text and 'exact match' in hdr_text:
        # Col 0 (Model Variant) wide enough for "Variant D (CORN): Conditional..."
        return [1.55, 0.85, 0.55, 0.75, 0.80, 0.65, 0.65, 0.60, 0.60], Pt(7.2), (50, 50)

    # 4. Table 3: Per-Class Diagnostic Performance (5 columns)
    if 'recall / sensitivity' in hdr_text:
        return [1.20, 1.80, 1.35, 1.30, 1.35], Pt(8.0), (70, 70)

    # 5. Table 4: Accuracy Enhancement (7 columns)
    if 'model configuration' in hdr_text and 'catastrophic' in hdr_text:
        return [1.70, 0.90, 0.85, 0.80, 0.90, 0.90, 0.95], Pt(7.5), (60, 60)

    # 6. Table 5: Post-Hoc Logit Adjustment Temperature Sweep (6 columns)
    if 'logit adjustment' in hdr_text:
        return [1.80, 0.85, 0.85, 0.85, 1.30, 1.35], Pt(8.0), (70, 70)

    # Default fallback: uniform distribution across 7.0 inches
    w = round(7.0 / num_cols, 2)
    return [w] * num_cols, Pt(8.0), (60, 60)

def render_table(doc, table_lines):
    """Renders markdown table lines into a beautifully formatted docx.Table with anti-wrapping column widths."""
    if len(table_lines) < 2:
        return

    # Extract rows
    rows_data = []
    for line in table_lines:
        parts = [p.strip() for p in line.strip('|').split('|')]
        # Filter separator row (e.g. ---|:---:|---)
        if all(re.match(r'^:?-+:?$', p) for p in parts):
            continue
        rows_data.append(parts)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    # Determine column widths and font size
    header_row = rows_data[0]
    hdr_text = " ".join(header_row).lower()
    col_widths, font_size, (pad_tb, pad_lr) = get_table_column_widths(header_row, num_cols)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    for r_idx, row_data in enumerate(rows_data):
        row = table.rows[r_idx]
        is_header = (r_idx == 0)

        # Anti-split rule on table row
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            # Apply explicit column width
            w_in = col_widths[c_idx] if c_idx < len(col_widths) else 1.0
            set_cell_width(cell, w_in)
            set_cell_margins(cell, top=pad_tb, bottom=pad_tb, left=pad_lr, right=pad_lr)

            text = row_data[c_idx] if c_idx < len(row_data) else ""
            p = cell.paragraphs[0]
            if 'student name' in hdr_text or 'prn' in hdr_text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or c_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

            if is_header:
                set_cell_background(cell, "2B5C8F")
                add_formatted_runs(p, text, base_font_size=font_size, base_bold=True, base_color=RGBColor(255, 255, 255))
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F7F9FC")
                else:
                    set_cell_background(cell, "FFFFFF")
                add_formatted_runs(p, text, base_font_size=font_size, base_color=RGBColor(30, 30, 30))

    # Add space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(6)

if __name__ == "__main__":
    import sys
    md_in = sys.argv[1] if len(sys.argv) > 1 else "./submission/final_paper.md"
    docx_out = sys.argv[2] if len(sys.argv) > 2 else "./submission/final_paper.docx"
    convert_markdown_to_docx(md_in, docx_out)
